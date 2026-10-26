"""Document Service (ТЗ 3.2.7): Excel (.xlsx), PDF, экспорт для Гранд-Смета/АРПС (CSV)."""
from __future__ import annotations

import csv
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import ROOT, settings
from .db import db
from .models import utcnow
from .security import audit

EXPORT_DIR = ROOT / "data" / "exports"
FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
SUBHEADER_FILL = PatternFill("solid", fgColor="EAF0F6")
GREEN = PatternFill("solid", fgColor="D8F3DC")
AMBER = PatternFill("solid", fgColor="FFF1C1")
RED = PatternFill("solid", fgColor="FFD6D6")
THIN = Side(style="thin", color="C9D3DF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = ["№", "Код", "Наименование работ", "Ед.", "Кол-во", "Цена, ₽", "Коэфф.",
           "Сумма без НДС, ₽", "НДС, ₽", "Итого, ₽", "Confidence", "Статус"]


def _rows(estimate: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for idx, p in enumerate(estimate["positions"], 1):
        coeffs = ", ".join(f"{k}={v}" for k, v in p["coefficients"].items())
        rows.append([idx, p["catalog_id"], p["name"], p["unit"], str(p["quantity"]),
                     str(p["unit_price"]), coeffs, str(p["subtotal"]), str(p["vat_amount"]),
                     str(p["total"]), p["confidence"], p["status"]])
    return rows


def _confidence_fill(conf: float) -> PatternFill:
    if conf >= 0.85:
        return GREEN
    if conf >= 0.5:
        return AMBER
    return RED


def _register_fonts() -> tuple[str, str]:
    regular, bold = "Helvetica", "Helvetica-Bold"
    if (FONT_DIR / "DejaVuSans.ttf").exists():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        regular, bold = "DejaVu", "DejaVu-Bold"
    return regular, bold


def generate_xlsx(estimate: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.merge_cells("A1:L1")
    ws["A1"] = f"{settings.company_name} — Локальная смета № {estimate['id']}"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    meta = [
        f"Заказчик: {estimate['client'] or '—'}    Объект: {estimate['title']}",
        f"Версия: {estimate['version_no']} от {estimate['version_created_at']}    Статус: {estimate['status']}",
        f"Регион: {estimate['region']}    Trace ID: {estimate['trace_id']}    Контрольная сумма: {estimate['checksum'][:16]}…",
        "Все суммы рассчитаны детерминированным кодом (Decimal). Позиции ссылаются на price_catalog.",
    ]
    for i, line in enumerate(meta, start=2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
        c = ws.cell(row=i, column=1, value=line)
        c.font = Font(size=10, italic=i == 5)
        c.alignment = Alignment(vertical="center")
    header_row = 7
    for col, name in enumerate(COLUMNS, 1):
        c = ws.cell(row=header_row, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for r_off, row in enumerate(_rows(estimate), start=header_row + 1):
        for col, value in enumerate(row, 1):
            c = ws.cell(row=r_off, column=col, value=value)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=(col == 3))
            if col in (5, 6, 8, 9, 10):
                c.alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=r_off, column=11).fill = _confidence_fill(estimate["positions"][r_off - header_row - 1]["confidence"])
    total_row = header_row + len(estimate["positions"]) + 2
    t = estimate["totals"]
    totals_lines = [
        ("Сумма без НДС и наценок", t["subtotal"]),
        ("Накладные расходы", t["overhead"]),
        ("Сметная прибыль", t["profit"]),
        ("НДС", t["vat"]),
        ("ИТОГО, ₽", t["grand_total"]),
    ]
    for i, (label, value) in enumerate(totals_lines):
        r = total_row + i
        ws.cell(row=r, column=9, value=label).font = Font(bold=i == 4)
        c = ws.cell(row=r, column=10, value=value)
        c.font = Font(bold=True, size=12 if i == 4 else 11)
        c.alignment = Alignment(horizontal="right")
        if i == 4:
            c.fill = GREEN
        for col in (9, 10):
            ws.cell(row=r, column=col).border = BORDER
    ws.cell(row=total_row + len(totals_lines) + 1, column=1,
            value="Верификация независимым пересчётом: " + t["verification"]).font = Font(italic=True)
    ws.cell(row=total_row + len(totals_lines) + 2, column=1,
            value="ЧЕРНОВИК: финальную ответственность несёт утвердивший смету человек." if estimate["status"] != "approved"
            else "Смета утверждена. Подписи: сметчик / руководитель.").font = Font(bold=True, color="B00020" if estimate["status"] != "approved" else "1B6B3A")

    ws2 = wb.create_sheet("Замечания")
    ws2.append(["Уровень", "Код", "Сообщение", "Рекомендация"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for issue in estimate.get("issues", []):
        ws2.append([issue["level"], issue["code"], issue["message"], issue.get("suggestion", "")])
    ws3 = wb.create_sheet("Предложения")
    ws3.append(["Код", "Наименование", "Ед.", "Цена", "Причина"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for s in estimate.get("suggestions", []):
        ws3.append([s["catalog_id"], s["name"], s["unit"], s["price"], s["reason"]])
    widths = {"A": 6, "B": 12, "C": 48, "D": 8, "E": 10, "F": 12, "G": 14, "H": 18, "I": 14, "J": 14, "K": 12, "L": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A8"
    for sheet in (ws2, ws3):
        for i, w in enumerate([12, 26, 70, 42, 20], 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def generate_pdf(estimate: dict[str, Any], path: Path) -> None:
    regular, bold = _register_fonts()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleRU", parent=styles["Title"], fontName=bold, fontSize=15, alignment=TA_CENTER, spaceAfter=6)
    small = ParagraphStyle("SmallRU", parent=styles["Normal"], fontName=regular, fontSize=7.5, leading=9.5)
    small_b = ParagraphStyle("SmallRUB", parent=small, fontName=bold)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4),
                            rightMargin=10 * mm, leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
                            title=f"Смета {estimate['id']}", author=settings.company_name)
    story: list[Any] = [
        Paragraph(f"{settings.company_name} — смета № {estimate['id']}", title_style),
        Paragraph(f"Заказчик: {estimate['client'] or '—'} | Объект: {estimate['title']} | "
                  f"Версия: {estimate['version_no']} | Статус: {estimate['status']} | "
                  f"Trace: {estimate['trace_id']}", small),
        Spacer(1, 5 * mm),
    ]
    data: list[list[Any]] = [[Paragraph(h, small_b) for h in
                             ["№", "Код", "Наименование", "Ед.", "Кол-во", "Цена", "Сумма", "НДС", "Итого", "Conf.", "Статус"]]]
    for idx, p in enumerate(estimate["positions"], 1):
        data.append([Paragraph(str(idx), small), Paragraph(p["catalog_id"], small),
                     Paragraph(p["name"], small), Paragraph(p["unit"], small),
                     Paragraph(str(p["quantity"]), small), Paragraph(str(p["unit_price"]), small),
                     Paragraph(str(p["subtotal"]), small), Paragraph(str(p["vat_amount"]), small),
                     Paragraph(str(p["total"]), small), Paragraph(f"{p['confidence']:.2f}", small),
                     Paragraph(p["status"], small)])
    table = Table(data, colWidths=[9 * mm, 18 * mm, 86 * mm, 12 * mm, 17 * mm, 19 * mm, 22 * mm, 20 * mm, 22 * mm, 14 * mm, 24 * mm], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C4D2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for i, p in enumerate(estimate["positions"], 1):
        c = p["confidence"]
        color = colors.HexColor("#D8F3DC") if c >= 0.85 else (colors.HexColor("#FFF1C1") if c >= 0.5 else colors.HexColor("#FFD6D6"))
        style.append(("BACKGROUND", (9, i), (9, i), color))
    table.setStyle(TableStyle(style))
    story.append(table)
    t = estimate["totals"]
    story.append(Spacer(1, 4 * mm))
    totals_table = Table([
        [Paragraph("Сумма без НДС", small_b), Paragraph(t["subtotal"], small),
         Paragraph("Накладные", small_b), Paragraph(t["overhead"], small),
         Paragraph("Прибыль", small_b), Paragraph(t["profit"], small)],
        [Paragraph("НДС", small_b), Paragraph(t["vat"], small),
         Paragraph("ИТОГО, ₽", small_b), Paragraph(t["grand_total"], small_b),
         Paragraph("Верификация", small_b), Paragraph(t["verification"], small)],
    ], colWidths=[35 * mm, 35 * mm, 30 * mm, 35 * mm, 35 * mm, 35 * mm])
    totals_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C4D2")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF0F6")),
    ]))
    story.append(totals_table)
    story.append(Spacer(1, 3 * mm))
    for issue in estimate.get("issues", [])[:8]:
        story.append(Paragraph(f"• [{issue['level']}] {issue['message']}", small))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Документ сформирован локально. Цветовая маркировка: зелёный — confidence ≥ 0.85, "
                           "жёлтый — 0.5–0.85, красный — < 0.5 (обязательная ручная проверка).", small))
    story.append(Paragraph("ЧЕРНОВИК — финальную ответственность за смету несёт утвердивший её человек."
                           if estimate["status"] != "approved" else "Смета утверждена.", small_b))
    doc.build(story)


def generate_grand_csv(estimate: dict[str, Any], path: Path) -> None:
    """Упрощённый CSV для импорта в Гранд-Смета/АРПС (формат требует сверки с версией ПО)."""
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["№", "Обоснование", "Наименование работ", "Ед. изм.", "Кол-во", "Цена за ед., руб.", "Стоимость, руб."])
        for idx, p in enumerate(estimate["positions"], 1):
            writer.writerow([idx, p["catalog_id"], p["name"], p["unit"], p["quantity"],
                             p["unit_price"], p["total"]])
        writer.writerow(["", "", "ИТОГО", "", "", "", estimate["totals"]["grand_total"]])


def generate(estimate: dict[str, Any], fmt: str, actor: dict[str, Any]) -> dict[str, Any]:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    suffix = {"xlsx": ".xlsx", "pdf": ".pdf", "grand_smeta_csv": ".csv"}[fmt]
    path = EXPORT_DIR / f"{estimate['id']}_v{estimate['version_no']}{suffix}"
    if fmt == "xlsx":
        generate_xlsx(estimate, path)
    elif fmt == "pdf":
        generate_pdf(estimate, path)
    else:
        generate_grand_csv(estimate, path)
    db.insert("documents", {
        "id": str(uuid.uuid4()), "smeta_id": estimate["id"], "version_no": estimate["version_no"],
        "fmt": fmt, "path": str(path), "created_at": utcnow(), "created_by": actor.get("username", "anonymous"),
    })
    audit(actor, "generate_document", "document", estimate["id"], {"fmt": fmt, "version": estimate["version_no"]})
    return {"fmt": fmt, "filename": path.name, "path": str(path), "size_bytes": path.stat().st_size,
            "smeta_id": estimate["id"], "version_no": estimate["version_no"]}
